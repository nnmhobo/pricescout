"""Export routes: /api/export/excel, /api/export/price-matrix, /api/exports,
/api/exports/add-urls (URL enrichment of an uploaded кошторис)."""

import io
import re
import tempfile
from datetime import datetime
from pathlib import Path

import pandas as pd
from flask import Blueprint, jsonify, request, send_file

from core.core import state, EXPORTS_DIR, ensure_exports_dir

bp = Blueprint("exports", __name__)

# Exports are grouped by origin:
#   monitoring/ — Excel files produced by monitoring runs & price matrices
#   urls/       — uploaded кошторис files enriched with a URL column
# Legacy files that predate the split stay in the root ("root" type).
EXPORT_DIRS = {
    "monitoring": EXPORTS_DIR / "monitoring",
    "urls":       EXPORTS_DIR / "urls",
    "root":       EXPORTS_DIR,
}


def ensure_export_dirs():
    ensure_exports_dir()
    EXPORT_DIRS["monitoring"].mkdir(parents=True, exist_ok=True)
    EXPORT_DIRS["urls"].mkdir(parents=True, exist_ok=True)

COLUMN_ORDER = ["item_label",
                "name", "price", "currency", "unit",
                "qty", "unit_estimate", "total_price",
                "brand", "sku",
                "specs", "supplier", "comment", "url", "date_scraped"]

COLUMN_NAMES = {
    "item_label":    "Матеріал кошторису",
    "name":          "Назва товару",
    "price":         "Ціна за од.",
    "currency":      "Валюта",
    "unit":          "Од.",
    "qty":           "К-сть (кошторис)",
    "unit_estimate": "Од. (кошторис)",
    "total_price":   "Загальна ціна",
    "brand":         "Бренд",
    "sku":           "Артикул",
    "specs":         "Характеристики",
    "supplier":      "Постачальник",
    "comment":       "Коментар",
    "url":           "Посилання",
    "date_scraped":  "Дата",
}


# Gold / silver / bronze fills for top-3 prices within each `item_label`
# group. The leading colour matches the in-app UI medal palette.
_MEDAL_FILLS = {
    1: ("FFF4C7", "8A6B0A"),  # gold-ish background, dark amber text
    2: ("E5E5E5", "454c54"),  # silver
    3: ("F4DBC1", "8a4a17"),  # bronze
}


def _build_excel(results: list, label: str) -> tuple[bytes, str]:
    from openpyxl.styles import Font, PatternFill, Alignment

    # Compute the per-group price rank BEFORE we drop columns / rename —
    # we still have `item_label` and `price` in the original list shape.
    # `ranks` is parallel to `results`: ranks[i] is 1/2/3 or None.
    grouped_prices: dict[str, list[float]] = {}
    for r in results:
        p = r.get("price")
        # Skip missing AND zero prices: fill_missing stub rows carry price 0
        # and must never win a medal.
        if p is None or float(p) <= 0:
            continue
        key = r.get("item_label") or r.get("name") or ""
        grouped_prices.setdefault(key, []).append(float(p))
    # Distinct sorted prices per group → 1/2/3 lookup; tied prices share a rank.
    group_rank_table: dict[str, dict[float, int]] = {}
    for key, prices in grouped_prices.items():
        ordered = sorted(set(prices))
        group_rank_table[key] = {p: i + 1 for i, p in enumerate(ordered[:3])}

    ranks: list[int | None] = []
    for r in results:
        p = r.get("price")
        if p is None or float(p) <= 0:
            ranks.append(None)
            continue
        key = r.get("item_label") or r.get("name") or ""
        ranks.append(group_rank_table.get(key, {}).get(float(p)))

    df = pd.DataFrame(results)
    if "specs" in df.columns:
        df["specs"] = df["specs"].apply(
            lambda x: ", ".join(f"{k}: {v}" for k, v in x.items())
            if isinstance(x, dict) else ""
        )
    df = df[[c for c in COLUMN_ORDER if c in df.columns]]
    df.rename(columns=COLUMN_NAMES, inplace=True)

    safe_lbl = re.sub(r'[^\w\-]', '_', label)[:40]
    filename  = f"{safe_lbl}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Ціни")
        ws = writer.sheets["Ціни"]
        for col in ws.columns:
            w = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(w + 4, 60)
        fill = PatternFill("solid", fgColor="141820")
        font = Font(bold=True, color="FFFFFF", size=11)
        for cell in ws[1]:
            cell.fill      = fill
            cell.font      = font
            cell.alignment = Alignment(horizontal="center")

        # Apply the medal fill to the entire row for top-3 entries in each
        # group. Row offset = 2 because Excel is 1-indexed and row 1 is
        # the header. `price_col_idx` is the column index for "Ціна за од."
        # which gets the bold medal text colour on top of the row fill.
        header_row = [c.value for c in ws[1]]
        try:
            price_col_idx = header_row.index(COLUMN_NAMES["price"]) + 1
        except ValueError:
            price_col_idx = None
        for row_idx, rank in enumerate(ranks, start=2):
            if rank is None or rank not in _MEDAL_FILLS:
                continue
            bg, fg = _MEDAL_FILLS[rank]
            row_fill = PatternFill("solid", fgColor=bg)
            for cell in ws[row_idx]:
                cell.fill = row_fill
            if price_col_idx is not None:
                ws.cell(row=row_idx, column=price_col_idx).font = Font(bold=True, color=fg)
    return output.getvalue(), filename


@bp.route("/api/export/excel")
def export_excel():
    if not state["results"]:
        return jsonify({"error": "Немає даних для експорту"}), 400
    excel_bytes, filename = _build_excel(state["results"], state.get("label", "prices"))
    ensure_export_dirs()
    (EXPORT_DIRS["monitoring"] / filename).write_bytes(excel_bytes)
    return send_file(
        io.BytesIO(excel_bytes),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


@bp.route("/api/exports")
def list_exports():
    ensure_export_dirs()
    out = []
    for etype in ("monitoring", "urls", "root"):
        d = EXPORT_DIRS[etype]
        for f in d.glob("*.xlsx"):
            if etype == "root" and f.parent != EXPORTS_DIR:
                continue
            out.append({
                "filename": f.name,
                "type":     etype,
                "size_kb":  round(f.stat().st_size / 1024, 1),
                "mtime":    f.stat().st_mtime,
                "created":  datetime.fromtimestamp(f.stat().st_mtime).strftime("%d.%m.%Y %H:%M"),
            })
    out.sort(key=lambda x: x["mtime"], reverse=True)
    for o in out:
        o.pop("mtime", None)
    return jsonify(out)


def _safe_export_path(filename: str, etype: str = "root"):
    """Resolve a user-supplied filename inside one of the export dirs.
    Returns None if the type is unknown or the result would escape the
    directory (e.g. via '..')."""
    base_dir = EXPORT_DIRS.get(etype)
    if base_dir is None:
        return None
    safe = re.sub(r'[^\w\-\.]', '', filename)
    if not safe or safe in (".", ".."):
        return None
    ensure_export_dirs()
    base = base_dir.resolve()
    candidate = (base / safe).resolve()
    try:
        candidate.relative_to(base)
    except ValueError:
        return None
    return candidate


@bp.route("/api/exports/<filename>")
def download_export(filename):
    # Legacy route — files that predate the monitoring/ + urls/ split.
    path = _safe_export_path(filename)
    if path is None or not path.exists():
        return jsonify({"error": "Файл не знайдено"}), 404
    return send_file(path, as_attachment=True, download_name=path.name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.route("/api/exports/<filename>", methods=["DELETE"])
def delete_export(filename):
    path = _safe_export_path(filename)
    if path is not None and path.exists():
        path.unlink()
    return jsonify({"status": "deleted"})


@bp.route("/api/exports/<etype>/<filename>")
def download_export_typed(etype, filename):
    path = _safe_export_path(filename, etype)
    if path is None or not path.exists():
        return jsonify({"error": "Файл не знайдено"}), 404
    return send_file(path, as_attachment=True, download_name=path.name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.route("/api/exports/<etype>/<filename>", methods=["DELETE"])
def delete_export_typed(etype, filename):
    path = _safe_export_path(filename, etype)
    if path is not None and path.exists():
        path.unlink()
    return jsonify({"status": "deleted"})


@bp.route("/api/export/price-matrix")
def export_price_matrix():
    """Export all monitored items × all suppliers as a price matrix Excel."""
    from core.item_db import get_availability_matrix
    from core.suppliers import SUPPLIERS
    from openpyxl.styles import Font, PatternFill, Alignment

    sup_ids   = [s["id"]   for s in SUPPLIERS if s.get("enabled", True)]
    sup_names = {s["id"]: s["name"] for s in SUPPLIERS}

    items = get_availability_matrix()
    if not items:
        return jsonify({"error": "Немає даних — спочатку запустіть discovery"}), 400

    rows = []
    for item in items:
        sups = item.get("suppliers", {})
        prices_found = []
        row = {
            "Матеріал":  item["label"],
            "Категорія": item.get("category") or "",
        }
        for sid in sup_ids:
            entry = sups.get(sid)
            price = None
            if entry and entry.get("found") and entry.get("last_price"):
                price = entry["last_price"]
                prices_found.append(price)
            row[sup_names.get(sid, sid)] = price

        row["Мін. ціна ₴"]  = min(prices_found) if prices_found else None
        row["Макс. ціна ₴"] = max(prices_found) if prices_found else None
        row["Знайдено (сайтів)"] = len(prices_found)
        if len(prices_found) >= 2:
            mn, mx = min(prices_found), max(prices_found)
            row["Розкид %"] = round((mx - mn) / mn * 100, 1) if mn > 0 else None
        else:
            row["Розкид %"] = None
        rows.append(row)

    df = pd.DataFrame(rows)
    # Sort by category then name
    df = df.sort_values(["Категорія", "Матеріал"])

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Прайс-матриця")
        ws = writer.sheets["Прайс-матриця"]

        header_fill = PatternFill("solid", fgColor="141820")
        header_font = Font(bold=True, color="FFFFFF", size=10)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Auto column widths
        for col_cells in ws.columns:
            width = max(len(str(c.value or "")) for c in col_cells)
            ws.column_dimensions[col_cells[0].column_letter].width = min(width + 3, 40)

        # Highlight min price cells (green) in supplier columns
        sup_col_indices = {}
        for cell in ws[1]:
            if cell.value in sup_names.values():
                sup_col_indices[cell.column] = True

        for row_num in range(2, ws.max_row + 1):
            # Find min value in supplier columns for this row
            sup_vals = []
            for col_idx in sup_col_indices:
                v = ws.cell(row=row_num, column=col_idx).value
                if isinstance(v, (int, float)):
                    sup_vals.append((v, col_idx))
            if sup_vals:
                min_val = min(v for v, _ in sup_vals)
                for v, col_idx in sup_vals:
                    cell = ws.cell(row=row_num, column=col_idx)
                    if v == min_val:
                        cell.font = Font(color="1a6b5a", bold=True)
                        cell.fill = PatternFill("solid", fgColor="e6f4f1")

    output.seek(0)
    ensure_export_dirs()
    filename = f"price_matrix_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    (EXPORT_DIRS["monitoring"] / filename).write_bytes(output.getvalue())
    output.seek(0)
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


# ── URL enrichment of an uploaded кошторис ─────────────────────
# The customer uploads their АВК-5 Excel; rows are matched to DB items by
# АВК code (fallback: exact material name) and a column with the known URL
# (cheapest supplier's product page) is written in. The customer's file is
# preserved: .xlsx keeps its formatting; .xls is converted to .xlsx
# values-only (the old format can't be written back).

_VARIANT_RE = re.compile(r'\s*варіант\s*\d+', re.IGNORECASE)


def _clean_code(v) -> str:
    if v is None:
        return ""
    return _VARIANT_RE.sub('', str(v).replace('\n', ' ').strip()).strip()


def _clean_name(v) -> str:
    if v is None:
        return ""
    s = str(v).replace('\n', ' ').strip()
    return "" if s == 'nan' else s


def _best_url_for_item(item: dict) -> str:
    """Cheapest found supplier's URL (manual price honoured); else any
    found URL; else ''."""
    best_url, best_price, any_url = "", None, ""
    for e in (item.get("suppliers") or {}).values():
        if not e.get("found"):
            continue
        url = e.get("url") or ""
        if not url:
            continue
        if not any_url:
            any_url = url
        price = e.get("manual_price") if e.get("manual_price") is not None else e.get("last_price")
        if price is not None and (best_price is None or price < best_price):
            best_price, best_url = price, url
    return best_url or any_url


def _build_url_lookup() -> tuple[dict, dict]:
    """{avk_code: url}, {label_lower: url} for every item with a known URL."""
    from core.item_db import load_items
    by_code: dict = {}
    by_name: dict = {}
    for item in load_items():
        url = _best_url_for_item(item)
        if not url:
            continue
        code = (item.get("avk_code") or "").strip()
        if code:
            by_code.setdefault(code, url)
        label = (item.get("label") or "").strip().lower()
        if label:
            by_name.setdefault(label, url)
    return by_code, by_name


def _load_upload_as_workbook(tmp_path: Path, suffix: str):
    """Return (openpyxl.Workbook, converted_from_xls). .xlsx is loaded
    natively (formatting preserved); .xls is copied values-only."""
    from openpyxl import Workbook, load_workbook
    if suffix == ".xlsx":
        return load_workbook(str(tmp_path)), False
    import xlrd
    wb_old = xlrd.open_workbook(str(tmp_path))
    ws_old = wb_old.sheet_by_index(0)
    wb = Workbook()
    ws = wb.active
    for r in range(ws_old.nrows):
        for c in range(ws_old.ncols):
            v = ws_old.cell_value(r, c)
            if v != '':
                ws.cell(row=r + 1, column=c + 1, value=v)
    return wb, True


# 1-based column of the material name per кошторис layout. Code is column B
# in both; КД_РЛМТ has an extra «Варіант ціни» column C, so its name is in D.
_NAME_COL = {"pvr": 3, "rlmt": 4}


def _detect_layout(ws) -> str:
    """'rlmt' if column A has a «Розділ N» section marker, else 'pvr' —
    the same signal the two parsers use to tell the layouts apart."""
    from parsers.rlmt_parser import SECTION_RE
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is not None and SECTION_RE.match(str(v).strip()):
            return "rlmt"
    return "pvr"


def _enrich_worksheet(ws, by_code: dict, by_name: dict,
                      target_col: int | None, fill_all_dups: bool) -> dict:
    """Write URLs into `target_col` (or one past the last used column).
    Works for both КД_ПВР and КД_РЛМТ (layout auto-detected).
    Returns stats: rows (item rows), filled, dup_skipped, not_found,
    column, layout."""
    from parsers.kostoris_parser import CODE_RE
    layout = _detect_layout(ws)
    name_col = _NAME_COL[layout]
    col = target_col or (ws.max_column + 1)
    stats = {"rows": 0, "filled": 0, "dup_skipped": 0, "not_found": 0,
             "column": col, "layout": layout}
    seen_keys: set = set()
    header_done = False
    for r in range(1, ws.max_row + 1):
        code = _clean_code(ws.cell(row=r, column=2).value)
        name = _clean_name(ws.cell(row=r, column=name_col).value)
        if not header_done and ("шифр" in code.lower() or "найменування" in name.lower()):
            ws.cell(row=r, column=col, value="Посилання")
            header_done = True
            continue
        if not CODE_RE.match(code) or not name:
            continue
        stats["rows"] += 1
        url = by_code.get(code) or by_name.get(name.lower())
        if not url:
            stats["not_found"] += 1
            continue
        key = code or name.lower()
        if not fill_all_dups and key in seen_keys:
            stats["dup_skipped"] += 1
            continue
        seen_keys.add(key)
        ws.cell(row=r, column=col, value=url)
        stats["filled"] += 1
    return stats


@bp.route("/api/exports/add-urls", methods=["POST"])
def add_urls():
    """Body (multipart): file=.xls/.xlsx АВК-5 кошторис.
    Form fields:
      column — Excel column letter to write URLs into (e.g. "K");
               empty = append after the last used column (default).
      dups   — "all" (default): every repeat row of the same material gets
               the URL (1:1); "first": only the first occurrence.
    Saves the enriched file to exports/urls/ and returns stats + download URL.
    """
    if "file" not in request.files:
        return jsonify({"error": "Файл не завантажено"}), 400
    f = request.files["file"]
    if not f.filename:
        return jsonify({"error": "Порожній файл"}), 400
    suffix = Path(f.filename).suffix.lower()
    if suffix not in (".xls", ".xlsx"):
        return jsonify({"error": f"Непідтримуваний формат '{suffix}'. Потрібен .xls або .xlsx."}), 400

    f.seek(0, 2); size = f.tell(); f.seek(0)
    if size < 1000:
        return jsonify({"error": "Файл занадто малий."}), 400
    if size > 20 * 1024 * 1024:
        return jsonify({"error": "Файл занадто великий (>20MB)."}), 400

    column_raw = (request.form.get("column") or "").strip()
    target_col = None
    if column_raw:
        if not re.fullmatch(r"[A-Za-z]{1,3}", column_raw):
            return jsonify({"error": "Невірна колонка — вкажіть літеру, напр. K"}), 400
        from openpyxl.utils import column_index_from_string
        target_col = column_index_from_string(column_raw.upper())
    fill_all_dups = (request.form.get("dups") or "all") != "first"

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
    tmp_path = Path(tmp.name)
    tmp.close()
    f.save(tmp_path)

    try:
        by_code, by_name = _build_url_lookup()
        wb, converted = _load_upload_as_workbook(tmp_path, suffix)
        stats = _enrich_worksheet(wb.active, by_code, by_name, target_col, fill_all_dups)

        stem = re.sub(r'[^\w\-]', '_', Path(f.filename).stem)[:40] or "kostoris"
        out_name = f"{stem}_посилання_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        ensure_export_dirs()
        wb.save(EXPORT_DIRS["urls"] / out_name)

        return jsonify({
            "filename":  out_name,
            "type":      "urls",
            "download":  f"/api/exports/urls/{out_name}",
            "converted_from_xls": converted,
            **stats,
        })
    except Exception as e:
        return jsonify({"error": f"Помилка обробки файлу: {e}"}), 500
    finally:
        tmp_path.unlink(missing_ok=True)