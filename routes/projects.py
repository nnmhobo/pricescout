"""Project management routes: /api/projects"""

import io
import re
from datetime import datetime

from flask import Blueprint, jsonify, request, send_file

from item_db import (
    create_project, get_projects, get_project, update_project,
    delete_project, get_project_items,
)

bp = Blueprint("projects", __name__)


def _enrich_items(items: list) -> list:
    """Add best_price, total_best, total_estimate, saving_pct to each item.

    Honours `manual_price` if the user corrected the parser's price by
    hand: the corrected value participates in best-price selection just
    like a freshly scraped one.
    """
    result = []
    for item in items:
        sups = item.get("suppliers", {})
        prices = [
            (v.get("manual_price") if v.get("manual_price") is not None else v["last_price"], sid)
            for sid, v in sups.items()
            if v.get("found") and (v.get("manual_price") is not None or v.get("last_price"))
        ]
        prices.sort(key=lambda x: x[0])
        best_price = prices[0][0] if prices else None
        best_supplier = prices[0][1] if prices else None
        qty = item.get("qty")
        est_price = item.get("estimate_unit_price")
        total_best = round(best_price * qty, 2) if best_price and qty else None
        total_est  = round(est_price * qty, 2)  if est_price  and qty else None
        saving_pct = None
        if est_price and best_price and est_price > 0:
            saving_pct = round((est_price - best_price) / est_price * 100, 1)
        result.append({
            **item,
            "best_price":    best_price,
            "best_supplier": best_supplier,
            "total_best":    total_best,
            "total_estimate": total_est,
            "saving_pct":    saving_pct,
        })
    return result


@bp.route("/api/projects", methods=["GET"])
def list_projects():
    return jsonify(get_projects())


@bp.route("/api/projects", methods=["POST"])
def create_proj():
    data = request.json or {}
    name = (data.get("name") or "").strip()
    if not name:
        return jsonify({"error": "Назва не може бути порожньою"}), 400
    return jsonify(create_project(name, data.get("description"), data.get("avk_file")))


@bp.route("/api/projects/<project_id>", methods=["GET"])
def get_proj(project_id):
    project = get_project(project_id)
    if not project:
        return jsonify({"error": "Не знайдено"}), 404
    return jsonify(project)


@bp.route("/api/projects/<project_id>", methods=["PATCH"])
def patch_proj(project_id):
    data = request.json or {}
    update_project(project_id, data)
    return jsonify(get_project(project_id))


@bp.route("/api/projects/<project_id>", methods=["DELETE"])
def delete_proj(project_id):
    delete_project(project_id)
    return jsonify({"status": "deleted"})


@bp.route("/api/projects/<project_id>/items", methods=["GET"])
def project_items_route(project_id):
    items = get_project_items(project_id)
    return jsonify(_enrich_items(items))


@bp.route("/api/projects/<project_id>/summary", methods=["GET"])
def project_summary(project_id):
    project = get_project(project_id)
    if not project:
        return jsonify({"error": "Не знайдено"}), 404
    items = _enrich_items(get_project_items(project_id))

    total_estimate = sum(i["total_estimate"] or 0 for i in items)
    total_best     = sum(i["total_best"]     or 0 for i in items)
    items_with_price = sum(1 for i in items if i["best_price"])
    saving = total_estimate - total_best if total_estimate and total_best else 0
    saving_pct = round(saving / total_estimate * 100, 1) if total_estimate > 0 else 0

    return jsonify({
        "project":          project,
        "item_count":       len(items),
        "items_with_price": items_with_price,
        "total_estimate":   round(total_estimate, 2),
        "total_best":       round(total_best, 2),
        "saving":           round(saving, 2),
        "saving_pct":       saving_pct,
    })


@bp.route("/api/projects/<project_id>/export", methods=["GET"])
def export_project(project_id):
    """Export project items to Excel with estimate vs best-price comparison."""
    import pandas as pd
    from openpyxl.styles import Font, PatternFill, Alignment
    from suppliers import SUPPLIERS

    project = get_project(project_id)
    if not project:
        return jsonify({"error": "Не знайдено"}), 404

    items = _enrich_items(get_project_items(project_id))
    sup_ids = [s["id"] for s in SUPPLIERS if s.get("enabled", True)]
    sup_names = {s["id"]: s["name"] for s in SUPPLIERS}

    rows = []
    for item in items:
        sups = item.get("suppliers", {})
        row = {
            "Код АВК-5":        item.get("avk_code") or "",
            "Матеріал":         item.get("label") or "",
            "Категорія":        item.get("category") or "",
            "Од.":              item.get("unit") or "",
            "К-сть":            item.get("qty"),
            "Ціна кошт. ₴":    item.get("estimate_unit_price"),
            "Сума кошт. ₴":    item.get("total_estimate"),
        }
        for sid in sup_ids:
            entry = sups.get(sid)
            if entry and entry.get("found") and entry.get("last_price"):
                row[sup_names.get(sid, sid)] = entry["last_price"]
            else:
                row[sup_names.get(sid, sid)] = None
        row["Мін. ціна ₴"]   = item.get("best_price")
        row["Сума мін. ₴"]   = item.get("total_best")
        row["Економія %"]     = item.get("saving_pct")
        rows.append(row)

    df = pd.DataFrame(rows)

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Кошторис")
        ws = writer.sheets["Кошторис"]

        # Header style
        header_fill = PatternFill("solid", fgColor="1a2b3c")
        header_font = Font(bold=True, color="FFFFFF", size=10)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Column widths
        for col_cells in ws.columns:
            width = max(len(str(c.value or "")) for c in col_cells)
            ws.column_dimensions[col_cells[0].column_letter].width = min(width + 4, 45)

        # Highlight saving_pct column: green if positive
        saving_col = None
        for cell in ws[1]:
            if cell.value == "Економія %":
                saving_col = cell.column
                break
        if saving_col:
            for row_num in range(2, ws.max_row + 1):
                cell = ws.cell(row=row_num, column=saving_col)
                val = cell.value
                if isinstance(val, (int, float)) and val > 0:
                    cell.font = Font(color="1a6b5a", bold=True)
                elif isinstance(val, (int, float)) and val < 0:
                    cell.font = Font(color="c0392b")

    output.seek(0)
    safe_name = re.sub(r'[^\w\-]', '_', project["name"])[:40]
    filename = f"project_{safe_name}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )
