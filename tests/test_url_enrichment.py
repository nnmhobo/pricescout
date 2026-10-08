"""URL enrichment of an uploaded кошторис (routes/exports.py).

Regression: ``_enrich_worksheet`` used to read the material name from column C,
which in КД_РЛМТ files is «Варіант ціни» — so name matching silently failed
there and only code matches worked. The layout is now auto-detected.
"""

from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from openpyxl import Workbook

from routes.exports import _detect_layout, _enrich_worksheet

URL_A = "https://shop.example/a"
URL_B = "https://shop.example/b"
# Item "B" has no code match in the DB lookup — it can only match by name.
BY_CODE = {"С111-1": URL_A}
BY_NAME = {"цемент м500": URL_B}


def _pvr_sheet():
    wb = Workbook(); ws = wb.active
    ws.append(["№", "Шифр ресурсу", "Найменування", "Од.", "К-сть", None, "Ціна"])
    ws.append([1, "С111-1", "Плитка", "м2", 10, None, 100])
    ws.append([2, "С999-9", "Цемент М500", "т", 1, None, 900])
    return ws


def _rlmt_sheet():
    wb = Workbook(); ws = wb.active
    ws.append(["№", "Шифр матеріалу", "Варіант ціни", "Найменування", "Од.", "К-сть", "Ціна"])
    ws.append(["Розділ 1. Ціноутворюючі матеріали"])
    ws.append([1, "С111-1", 1, "Плитка", "м2", 10, 100])
    ws.append(["Разом:"])
    ws.append(["Розділ 2. Неціноутворюючі матеріали"])
    ws.append([2, "С999-9", 1, "Цемент М500", "т", 1, 900])
    return ws


def test_detects_layouts():
    assert _detect_layout(_pvr_sheet()) == "pvr"
    assert _detect_layout(_rlmt_sheet()) == "rlmt"


def test_pvr_matches_by_code_and_name():
    ws = _pvr_sheet()
    stats = _enrich_worksheet(ws, BY_CODE, BY_NAME, None, True)
    assert stats["layout"] == "pvr"
    assert (stats["rows"], stats["filled"], stats["not_found"]) == (2, 2, 0)
    col = stats["column"]
    assert ws.cell(row=2, column=col).value == URL_A
    assert ws.cell(row=3, column=col).value == URL_B


def test_rlmt_matches_by_name_from_column_d():
    ws = _rlmt_sheet()
    stats = _enrich_worksheet(ws, BY_CODE, BY_NAME, None, True)
    assert stats["layout"] == "rlmt"
    # section markers and «Разом:» rows are not item rows
    assert (stats["rows"], stats["filled"], stats["not_found"]) == (2, 2, 0)
    col = stats["column"]
    assert ws.cell(row=1, column=col).value == "Посилання"
    assert ws.cell(row=3, column=col).value == URL_A
    assert ws.cell(row=6, column=col).value == URL_B
